"""
IntegriGuard Main Backend — FastAPI
Fixes applied:
  - BUG-3: Full SQLAlchemy DB integration (Postgres in prod, SQLite in local dev)
  - BUG-4: Added GET /api/submissions, GET /api/submissions/{id}, POST /api/faculty-actions
  - BUG-5: Added GET /api/alerts, POST /api/alerts/{id}/read
  - Added:  GET /api/students, GET /api/students/{id}
  - Added:  POST /api/admin/* placeholders
  - Added:  POST /api/baseline-sample
"""
import os
import json
import uuid
import io
from datetime import datetime
from typing import Optional, List

import httpx
import docx
import pdfplumber
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Depends
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from sqlalchemy.orm import Session

from database import engine, get_db, Base
from models import (
    Submission, RiskScore, Alert, FacultyAction,
    Student, Course, User, BaselineSample
)

# ── Bootstrap DB tables (creates SQLite file on first run; no-op if Postgres already migrated) ──
Base.metadata.create_all(bind=engine)

import socketio

# ── Socket.io Setup ──
sio = socketio.AsyncServer(async_mode='asgi', cors_allowed_origins='*')
socket_app = socketio.ASGIApp(sio)

app = FastAPI(title="IntegriGuard Main Backend")

# Mount socketio
app.mount("/socket.io", socket_app)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

ML_SERVICE_URL = os.getenv("ML_SERVICE_URL", "http://localhost:8001")


# ══════════════════════════════════════════════════════════════════════
# Pydantic Schemas
# ══════════════════════════════════════════════════════════════════════

class TextSubmission(BaseModel):
    text: str
    course_id: str
    assignment_name: str
    student_id: Optional[str] = "ST-999"
    version_note: Optional[str] = None
    week: Optional[int] = 1
    current_grade: Optional[float] = None
    speed_delta: Optional[float] = None


class FacultyActionIn(BaseModel):
    submission_id: str
    faculty_id: Optional[str] = "FAC-001"
    action_type: str  # e.g. "reviewed", "dismissed", "escalated"


class BaselineSampleIn(BaseModel):
    studentId: str
    text: str
    timeTaken: Optional[int] = None


class CourseConfigIn(BaseModel):
    course_id: Optional[str] = None
    high_threshold: Optional[int] = 70
    medium_threshold: Optional[int] = 40
    alert_sensitivity: Optional[str] = "medium"


# ══════════════════════════════════════════════════════════════════════
# Internal helpers
# ══════════════════════════════════════════════════════════════════════

async def _call_ml(submission_id: str, text: str, student_id: str, course_id: str,
                   week: int = 1, current_grade: Optional[float] = None,
                   speed_delta: Optional[float] = None) -> Optional[dict]:
    """Calls the ML service and returns the parsed result, or None on failure."""
    try:
        async with httpx.AsyncClient(timeout=30.0) as client:
            resp = await client.post(
                f"{ML_SERVICE_URL}/api/v1/analyse",
                json={
                    "submission_id": submission_id,
                    "text": text,
                    "student_id": student_id,
                    "course_id": course_id,
                    "week": week,
                    "current_grade": current_grade,
                    "speed_delta": speed_delta,
                }
            )
            resp.raise_for_status()
            return resp.json()
    except Exception as e:
        print(f"[ML Service] unavailable: {e}")
        return None


def _persist_ml_result(
    db: Session,
    submission_id: str,
    student_id: str,
    analysis: dict,
):
    """Writes ML result to risk_scores and creates an alert if score is high."""
    risk_score_val = analysis.get("score")
    if risk_score_val is None:
        return  # Baseline sample, no score yet

    score_int = int(round(float(risk_score_val) * 100))  # Store as 0-100 int

    rs = RiskScore(
        id=str(uuid.uuid4()),
        submission_id=submission_id,
        score=score_int,
        ai_prob=analysis.get("ai_prob"),
        style_shift=analysis.get("style_shift"),
        grade_anomaly=analysis.get("grade_anomaly"),
        shap_explanation=json.dumps(analysis.get("shap_explanation", {})),
        baseline_corrupted=analysis.get("poisoned_baseline", False),
        created_at=datetime.utcnow(),
    )
    db.add(rs)

    # Auto-create alert for high-risk submissions (score >= 60)
    if score_int >= 60:
        level = "HIGH" if score_int >= 80 else "MEDIUM"
        alert_id = str(uuid.uuid4())
        alert = Alert(
            id=alert_id,
            student_id=student_id,
            submission_id=submission_id,
            signal_type=level,
            message=(
                f"{level} integrity risk detected (score: {score_int}/100). "
                f"AI probability: {round((analysis.get('ai_prob') or 0) * 100)}%. "
                f"Style shift: {round((analysis.get('style_shift') or 0) * 100)}%."
            ),
            read=False,
            created_at=datetime.utcnow(),
        )
        db.add(alert)
        
        # Real-time notification
        import asyncio
        asyncio.create_task(sio.emit('new_alert', {
            "id": alert_id,
            "student_id": student_id,
            "submission_id": submission_id,
            "signal_type": level,
            "message": alert.message,
            "created_at": alert.created_at.isoformat()
        }))

    db.commit()

    # Notify frontend that analysis is complete
    import asyncio
    asyncio.create_task(sio.emit('score_updated', {"submissionId": submission_id}))


# ══════════════════════════════════════════════════════════════════════
# Health
# ══════════════════════════════════════════════════════════════════════

@app.get("/")
async def root():
    return {"status": "IntegriGuard Backend is running"}


# ══════════════════════════════════════════════════════════════════════
# Submissions (BUG-4 fixes)
# ══════════════════════════════════════════════════════════════════════

@app.get("/api/submissions")
async def get_all_submissions(db: Session = Depends(get_db)):
    """Returns all submissions with their risk scores."""
    submissions = db.query(Submission).order_by(Submission.submitted_at.desc()).limit(200).all()
    result = []
    for s in submissions:
        rs = db.query(RiskScore).filter(RiskScore.submission_id == s.id).first()
        result.append({
            "id": s.id,
            "student_id": s.student_id,
            "course_id": s.course_id,
            "assignment_name": s.assignment_name,
            "word_count": s.word_count,
            "status": s.status,
            "submitted_at": s.submitted_at.isoformat() if s.submitted_at else None,
            "risk_score": rs.score if rs else None,
            "ai_prob": rs.ai_prob if rs else None,
        })
    return result


@app.get("/api/submissions/{submission_id}")
async def get_submission(submission_id: str, db: Session = Depends(get_db)):
    """Returns a single submission with full risk details."""
    s = db.query(Submission).filter(Submission.id == submission_id).first()
    if not s:
        raise HTTPException(status_code=404, detail="Submission not found")
    rs = db.query(RiskScore).filter(RiskScore.submission_id == s.id).first()
    return {
        "id": s.id,
        "student_id": s.student_id,
        "course_id": s.course_id,
        "assignment_name": s.assignment_name,
        "word_count": s.word_count,
        "status": s.status,
        "submitted_at": s.submitted_at.isoformat() if s.submitted_at else None,
        "risk_score": rs.score if rs else None,
        "ai_prob": rs.ai_prob if rs else None,
        "style_shift": rs.style_shift if rs else None,
        "grade_anomaly": rs.grade_anomaly if rs else None,
        "shap_explanation": json.loads(rs.shap_explanation) if (rs and rs.shap_explanation) else {},
        "baseline_corrupted": rs.baseline_corrupted if rs else False,
    }


@app.post("/api/submissions/text")
async def submit_text(submission: TextSubmission, db: Session = Depends(get_db)):
    submission_id = str(uuid.uuid4())
    word_count = len(submission.text.split())

    # Persist submission
    db_sub = Submission(
        id=submission_id,
        student_id=submission.student_id,
        course_id=submission.course_id,
        assignment_name=submission.assignment_name,
        version_note=submission.version_note,
        content_text=submission.text,
        word_count=word_count,
        submitted_at=datetime.utcnow(),
        status="processing",
    )
    db.add(db_sub)
    db.commit()

    # Call ML service
    analysis = await _call_ml(
        submission_id, submission.text, submission.student_id,
        submission.course_id, submission.week,
        submission.current_grade, submission.speed_delta
    )
    if analysis:
        _persist_ml_result(db, submission_id, submission.student_id, analysis)
        db_sub.status = "analysed"
    else:
        db_sub.status = "ml_unavailable"
    db.commit()

    return {"submission_id": submission_id, "status": db_sub.status}


@app.post("/api/submissions/upload")
async def upload_file(
    file: UploadFile = File(...),
    course_id: str = Form(...),
    assignment_name: str = Form(...),
    student_id: str = Form("ST-999"),
    week: int = Form(1),
    version_note: Optional[str] = Form(None),
    db: Session = Depends(get_db),
):
    MAX_SIZE = 10 * 1024 * 1024
    content = await file.read()
    if len(content) > MAX_SIZE:
        raise HTTPException(status_code=400, detail="File too large (max 10MB)")

    text = ""
    filename = file.filename.lower()
    try:
        if filename.endswith(".docx"):
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join([para.text for para in doc.paragraphs])
        elif filename.endswith(".pdf"):
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                text = "\n".join([page.extract_text() or "" for page in pdf.pages])
        elif filename.endswith(".txt"):
            text = content.decode("utf-8")
        else:
            raise HTTPException(status_code=400, detail="Unsupported file format")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Text extraction failed: {str(e)}")

    if not text.strip():
        raise HTTPException(status_code=400, detail="Could not extract text from file")

    submission_id = str(uuid.uuid4())
    db_sub = Submission(
        id=submission_id,
        student_id=student_id,
        course_id=course_id,
        assignment_name=assignment_name,
        version_note=version_note,
        content_text=text,
        word_count=len(text.split()),
        submitted_at=datetime.utcnow(),
        status="processing",
    )
    db.add(db_sub)
    db.commit()

    analysis = await _call_ml(submission_id, text, student_id, course_id, week)
    if analysis:
        _persist_ml_result(db, submission_id, student_id, analysis)
        db_sub.status = "analysed"
    else:
        db_sub.status = "ml_unavailable"
    db.commit()

    return {"submission_id": submission_id, "status": db_sub.status}


@app.get("/api/submissions/student/{student_id}")
async def get_student_submissions(student_id: str, db: Session = Depends(get_db)):
    """Returns all submissions for a specific student."""
    subs = db.query(Submission).filter(Submission.student_id == student_id).order_by(
        Submission.submitted_at.desc()
    ).all()

    if not subs:
        return []

    # Calculate versions per assignment name
    assignment_counts = {}
    result = []
    
    # We process in chronological order to assign v1, v2, etc. correctly
    for s in sorted(subs, key=lambda x: x.submitted_at):
        title = s.assignment_name or "Untitled"
        assignment_counts[title] = assignment_counts.get(title, 0) + 1
        
        result.append({
            "id": s.id,
            "title": title,
            "version": f"v{assignment_counts[title]}",
            "words": s.word_count,
            "time": s.submitted_at.isoformat(),
            "note": s.version_note
        })

    # Return in reverse chronological order for the frontend feed
    return list(reversed(result))


# ══════════════════════════════════════════════════════════════════════
# Faculty Actions (BUG-4 fix)
# ══════════════════════════════════════════════════════════════════════

@app.post("/api/faculty-actions")
async def log_faculty_action(action: FacultyActionIn, db: Session = Depends(get_db)):
    fa = FacultyAction(
        id=str(uuid.uuid4()),
        submission_id=action.submission_id,
        faculty_id=action.faculty_id,
        action_type=action.action_type,
        created_at=datetime.utcnow(),
    )
    db.add(fa)
    db.commit()
    return {"status": "logged", "action_id": fa.id}


# ══════════════════════════════════════════════════════════════════════
# Alerts (BUG-5 fix)
# ══════════════════════════════════════════════════════════════════════

@app.get("/api/alerts")
async def get_alerts(db: Session = Depends(get_db)):
    alerts = db.query(Alert).order_by(Alert.created_at.desc()).limit(100).all()
    return [
        {
            "id": a.id,
            "student_id": a.student_id,
            "submission_id": a.submission_id,
            "signal_type": a.signal_type,
            "message": a.message,
            "read": a.read,
            "created_at": a.created_at.isoformat() if a.created_at else None,
        }
        for a in alerts
    ]


@app.post("/api/alerts/{alert_id}/read")
async def mark_alert_read(alert_id: str, db: Session = Depends(get_db)):
    alert = db.query(Alert).filter(Alert.id == alert_id).first()
    if not alert:
        raise HTTPException(status_code=404, detail="Alert not found")
    alert.read = True
    db.commit()
    return {"status": "marked_read", "alert_id": alert_id}


# ══════════════════════════════════════════════════════════════════════
# Students
# ══════════════════════════════════════════════════════════════════════

@app.get("/api/students")
async def get_students(db: Session = Depends(get_db)):
    students = db.query(Student).all()
    return [
        {
            "id": s.id,
            "student_code": s.student_code,
            "full_name": s.full_name,
            "course_id": s.course_id,
            "year": s.year,
            "enrolment_status": s.enrolment_status,
            "is_esl": s.is_esl,
            "is_mature": s.is_mature,
        }
        for s in students
    ]


@app.get("/api/students/{student_id}")
async def get_student(student_id: str, db: Session = Depends(get_db)):
    s = db.query(Student).filter(Student.id == student_id).first()
    if not s:
        raise HTTPException(status_code=404, detail="Student not found")
    return {
        "id": s.id,
        "student_code": s.student_code,
        "full_name": s.full_name,
        "course_id": s.course_id,
        "year": s.year,
        "enrolment_status": s.enrolment_status,
        "is_esl": s.is_esl,
        "is_mature": s.is_mature,
    }


# ══════════════════════════════════════════════════════════════════════
# Baseline Sample
# ══════════════════════════════════════════════════════════════════════

@app.post("/api/baseline-sample")
async def submit_baseline_sample(sample: BaselineSampleIn, db: Session = Depends(get_db)):
    submission_id = str(uuid.uuid4())
    db_sub = Submission(
        id=submission_id,
        student_id=sample.studentId,
        course_id="BASELINE",
        assignment_name="Baseline Sample",
        content_text=sample.text,
        word_count=len(sample.text.split()),
        submitted_at=datetime.utcnow(),
        submission_duration_seconds=sample.timeTaken,
        status="baseline",
    )
    db.add(db_sub)

    bs = BaselineSample(
        id=str(uuid.uuid4()),
        student_id=sample.studentId,
        submission_id=submission_id,
        week_number=1,
        created_at=datetime.utcnow(),
    )
    db.add(bs)
    db.commit()
    return {"status": "baseline_stored", "submission_id": submission_id}


# ══════════════════════════════════════════════════════════════════════
# Admin endpoints (BUG-5 fix)
# ══════════════════════════════════════════════════════════════════════

@app.get("/api/admin/bias-audit")
async def bias_audit(db: Session = Depends(get_db)):
    """Returns aggregate bias audit metrics."""
    total = db.query(RiskScore).count()
    high_risk = db.query(RiskScore).filter(RiskScore.score >= 70).count()
    return {
        "total_submissions_analysed": total,
        "high_risk_count": high_risk,
        "false_positive_rate_estimate": 0.05,  # Placeholder
        "bias_flags": [],
    }


@app.post("/api/admin/recalibrate")
async def recalibrate():
    """Triggers model recalibration (placeholder for scheduled job)."""
    return {"status": "recalibration_queued", "estimated_duration_seconds": 120}


@app.post("/api/admin/course-config")
async def save_course_config(config: CourseConfigIn, db: Session = Depends(get_db)):
    """Saves per-course threshold configuration."""
    if config.course_id:
        course = db.query(Course).filter(Course.id == config.course_id).first()
        if course:
            course.high_threshold = config.high_threshold
            course.medium_threshold = config.medium_threshold
            course.alert_sensitivity = config.alert_sensitivity
            db.commit()
    return {"status": "saved", "config": config.dict()}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
